import os
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import report_store as store
import login
import main
from filelock import FileLock, Timeout

HEADER = ("Código PIN", "Fecha", "Hora", "Empresa")


def write_report(path, rows, header=HEADER):
    workbook = Workbook()
    workbook.active.title = "Canjes"
    workbook.active.append(header)
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)
    workbook.close()
    return path


def read_rows(path):
    with store.report_rows(path) as (_, rows):
        return list(rows)


class IncrementalReportsTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.data = self.root / "data"
        self.data.mkdir()
        self.delta = self.root / "delta.xlsx"
        self.legacy = self.data / "canjes-institucion_04-09-2026_12-00-00.xlsx"
        self.old = ("001", "03/09/2026", "12:00:00", "Company")
        self.overlap = ("002", "04/09/2026", "12:00:00", "Company")
        self.new = ("003", "05/09/2026", "12:00:00", "Company")
        write_report(self.legacy, [self.old, self.overlap])

    def merge(self, baseline=None):
        return store.append_report(baseline or self.legacy, self.delta, self.data,
                                   date(2026, 9, 4), date(2026, 9, 5))

    def test_migration_preserves_history_appends_overlap_once_and_cleans_snapshots(self):
        older = self.data / "canjes-institucion_03-09-2026_12-00-00.xlsx"
        write_report(older, [self.old])
        unrelated = write_report(self.data / "Venta_Septiembre.xlsx", [])
        seed = write_report(self.root / self.legacy.name, [self.old])
        write_report(self.delta, [self.overlap, self.new])
        result = self.merge()
        self.assertEqual(Path(result).name, store.REPORT_FILENAME)
        self.assertEqual(read_rows(result), [self.old, self.overlap, self.new])
        self.assertEqual(store.checkpoint_day(result), date(2026, 9, 5))
        self.assertEqual(store.legacy_reports(self.data), [])
        self.assertTrue(unrelated.exists())
        self.assertTrue(seed.exists())
        workbook = load_workbook(result, read_only=True)
        self.assertEqual(workbook.worksheets[0].title, "Canjes")
        self.assertEqual(workbook[store.CHECKPOINT_SHEET].sheet_state, "hidden")
        workbook.close()

    def test_retries_are_idempotent_without_collapsing_identical_redemptions(self):
        write_report(self.legacy, [self.old, self.overlap, self.overlap])
        write_report(self.delta, [self.overlap, self.overlap, self.overlap, self.new])
        result = self.merge()
        self.assertEqual(read_rows(result), [self.old] + [self.overlap] * 3 + [self.new])
        self.merge(result)
        self.assertEqual(len(read_rows(result)), 5)

    def test_empty_delta_advances_checkpoint_without_losing_history(self):
        write_report(self.delta, [])
        result = self.merge()
        self.assertEqual(read_rows(result), [self.old, self.overlap])
        self.assertEqual(store.checkpoint_day(result), date(2026, 9, 5))

    def test_first_run_and_empty_seed_start_at_campaign_start(self):
        self.assertEqual(store.checkpoint_day(None), date(2026, 5, 7))
        write_report(self.delta, [])
        self.assertEqual(store.checkpoint_day(self.delta), date(2026, 5, 7))
        result = store.append_report(None, self.delta, self.data,
                                     date(2026, 5, 7), date(2026, 9, 5))
        self.assertEqual(read_rows(result), [])

    def test_legacy_checkpoint_uses_data_not_download_timestamp(self):
        self.assertEqual(store.checkpoint_day(self.legacy), date(2026, 9, 4))
        newer = self.data / "canjes-institucion_05-09-2026_12-00-00.xlsx"
        write_report(newer, [self.old])
        self.assertEqual(store.existing_report(self.data, self.root), newer)
        self.assertEqual(store.checkpoint_day(newer), date(2026, 9, 3))

    def test_invalid_delta_preserves_baseline(self):
        original = self.legacy.read_bytes()
        cases = [([self.old], HEADER), ([self.new], HEADER + ("Extra",)),
                 ([("003", "bad date", "12:00:00", "Company")], HEADER)]
        for rows, header in cases:
            with self.subTest(rows=rows, header=header):
                write_report(self.delta, rows, header)
                with self.assertRaises(ValueError):
                    self.merge()
                self.assertEqual(self.legacy.read_bytes(), original)
                self.assertFalse((self.data / store.REPORT_FILENAME).exists())

    def test_failed_publish_keeps_data_and_checkpoint_together(self):
        write_report(self.delta, [self.overlap])
        result = self.merge()
        original = Path(result).read_bytes()
        with patch.object(store.os, "replace", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                self.merge(result)
        self.assertEqual(Path(result).read_bytes(), original)
        self.assertEqual(list(self.data.glob(".canjes-merge-*")), [])

    def test_date_objects_and_text_dates_match_in_overlap(self):
        write_report(self.delta, [("002", datetime(2026, 9, 4), "12:00:00", "Company")])
        result = self.merge()
        self.assertEqual(len(read_rows(result)), 2)

    def test_refresh_requests_checkpoint_through_lima_today_and_cleans_temp_files(self):
        def download(start, end, directory, timeout):
            self.assertEqual(start, date(2026, 9, 4))
            self.assertEqual(end, date(2026, 9, 5))
            self.assertEqual(timeout, 99)
            return str(write_report(Path(directory) / "export.xlsx", [self.overlap, self.new]))

        clock = Mock()
        clock.now.return_value = datetime(2026, 9, 5, 12, tzinfo=store.LIMA_TZ)
        with patch.dict(os.environ, {"DATA_DIR": str(self.data)}), \
                patch.object(login, "datetime", clock), \
                patch.object(login, "_download_report", side_effect=download):
            result = login.run_report_extraction(99)
        clock.now.assert_called_once_with(store.LIMA_TZ)
        self.assertEqual(len(read_rows(result)), 3)
        self.assertEqual(list(self.data.glob(".canjes-download-*")), [])

    def test_failed_download_does_not_migrate_or_advance_checkpoint(self):
        def download(start, end, directory, timeout):
            (Path(directory) / "partial.crdownload").write_bytes(b"incomplete")
            return None
        with patch.dict(os.environ, {"DATA_DIR": str(self.data)}), \
                patch.object(login, "_download_report", side_effect=download):
            self.assertIsNone(login.run_report_extraction())
        self.assertTrue(self.legacy.exists())
        self.assertFalse((self.data / store.REPORT_FILENAME).exists())
        self.assertEqual(list(self.data.glob(".canjes-download-*")), [])

    def test_overlapping_refreshes_are_rejected(self):
        with patch.dict(os.environ, {"DATA_DIR": str(self.data)}), \
                FileLock(str(self.data / ".canjes-refresh.lock")), \
                patch.object(login, "_download_report") as download:
            with self.assertRaises(Timeout):
                login.run_report_extraction()
            download.assert_not_called()

    def test_latest_api_prefers_cumulative_file_and_uses_current_mtime(self):
        write_report(self.delta, [self.overlap, self.new])
        result = self.merge()
        write_report(self.data / "canjes-institucion_01-01-2099_12-00-00.xlsx", [])
        with patch.object(main, "DOWNLOAD_DIR", str(self.data)):
            payload = main.get_latest_report()
            self.assertEqual(payload["file_path"], result)
            self.assertTrue(payload["is_recent"])
            response = main.download_latest_report(username="test")
            self.assertEqual(response.path, result)

    def test_calendar_clicks_actual_checkpoint_and_today_including_same_day(self):
        for start in [date(2026, 8, 31), date(2026, 9, 5)]:
            driver = Mock()
            driver.find_elements.return_value = [Mock()]
            wait = Mock()
            with patch.object(login, "WebDriverWait", return_value=wait), \
                    patch.object(login.time, "sleep"):
                login.select_calendar_range(driver, start, date(2026, 9, 5))
            self.assertEqual(wait.until.call_count, 2)
            # Exercise Selenium's locator conditions to check the dates requested.
            for call in wait.until.call_args_list:
                call.args[0](driver)
            xpaths = [call.args[1] for call in driver.find_element.call_args_list]
            self.assertEqual(xpaths, [f"//button[@data-day='{start.day}/{start.month}/{start.year}']",
                                     "//button[@data-day='5/9/2026']"])

    def test_failed_calendar_selection_cancels_export_and_closes_browser(self):
        driver = Mock()
        username = Mock()
        username.get_attribute.return_value = "text"
        driver.find_elements.return_value = [username]
        wait = Mock()
        with patch.object(login, "_get_config_value", return_value="test"), \
                patch.object(login.shutil, "which", return_value="/test/chromedriver"), \
                patch.object(login.webdriver, "Chrome", return_value=driver), \
                patch.object(login, "WebDriverWait", return_value=wait), \
                patch.object(login, "select_calendar_range", side_effect=RuntimeError("missing day")), \
                patch.object(login.time, "sleep"):
            with self.assertRaisesRegex(RuntimeError, "export cancelled"):
                login._download_report(date(2026, 9, 4), date(2026, 9, 5), str(self.data))
        # Password, reports link, date picker only: never request the export button.
        self.assertEqual(wait.until.call_count, 3)
        driver.quit.assert_called_once()


if __name__ == "__main__":
    unittest.main()
