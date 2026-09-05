"""One cumulative workbook, with a checkpoint committed in the same XLSX file."""

import os
import re
from collections import Counter
from contextlib import contextmanager
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

LIMA_TZ = timezone(timedelta(hours=-5))
INITIAL_DATE = date(2026, 5, 7)
REPORT_FILENAME = "canjes-institucion.xlsx"
REPORT_NAME_RE = re.compile(
    r"canjes-institucion_(\d{2}-\d{2}-\d{4})_(\d{2}-\d{2}-\d{2})\.xlsx$",
    re.IGNORECASE,
)
CHECKPOINT_SHEET = "_checkpoint"


def data_directory():
    return Path(os.environ.get(
        "DATA_DIR", "/app/data" if os.environ.get("RENDER")
        else str(Path(__file__).resolve().parent / "downloads"),
    )).resolve()


def legacy_reports(*directories):
    return [p for directory in directories if Path(directory).is_dir()
            for p in Path(directory).iterdir()
            if p.is_file() and REPORT_NAME_RE.fullmatch(p.name)]


def existing_report(directory, project_dir):
    canonical = Path(directory) / REPORT_FILENAME
    if canonical.exists():
        return canonical
    candidates = legacy_reports(directory, project_dir)
    return max(candidates, key=lambda p: datetime.strptime(
        " ".join(REPORT_NAME_RE.fullmatch(p.name).groups()), "%d-%m-%Y %H-%M-%S"
    ), default=None)


def parse_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            pass
    raise ValueError(f"Invalid canjes Fecha: {value!r}")


@contextmanager
def report_rows(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = iter(sheet.values)
        header = next(rows, ())
        if "Fecha" not in header or "Código PIN" not in header:
            raise ValueError("The canjes workbook must contain Fecha and Código PIN columns.")
        yield header, (tuple(row) for row in rows if any(v is not None for v in row))
    finally:
        workbook.close()


def checkpoint_day(path):
    if path is None:
        return INITIAL_DATE
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if CHECKPOINT_SHEET in workbook.sheetnames:
            sheet = workbook[CHECKPOINT_SHEET]
            if sheet["A1"].value != "last_successful_day":
                raise ValueError("Invalid canjes checkpoint metadata.")
            return parse_date(sheet["B1"].value)
    finally:
        workbook.close()
    # Old full snapshots have no checkpoint. Re-fetch their last data day;
    # filename/mtime alone cannot prove that an entire day was downloaded.
    with report_rows(path) as (header, rows):
        index = header.index("Fecha")
        return max((parse_date(row[index]) for row in rows), default=INITIAL_DATE)


def _row_key(row, date_index):
    values = list(row)
    values[date_index] = parse_date(values[date_index]).isoformat()
    return tuple(values)


def append_report(baseline, incoming, directory, start_day, end_day):
    """Append unseen row occurrences, retaining legitimate identical canjes.

    Only the overlap window needs a Counter in memory. XLSX is a ZIP archive,
    so publishing the updated workbook requires a temporary rewrite, followed
    by atomic replacement. The data and checkpoint cannot diverge on failure.
    """
    target = Path(directory) / REPORT_FILENAME
    # Validate the complete delta BEFORE creating output or changing any files.
    with report_rows(incoming) as (header, rows):
        date_index = header.index("Fecha")
        for row in rows:
            if not start_day <= parse_date(row[date_index]) <= end_day:
                raise ValueError("Downloaded canjes fall outside the requested date range.")
    if baseline:
        with report_rows(baseline) as (old_header, _):
            if old_header != header:
                raise ValueError("Canjes columns changed; the existing workbook was preserved.")

    with TemporaryDirectory(prefix=".canjes-merge-", dir=directory) as temporary:
        staged = Path(temporary) / REPORT_FILENAME
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Canjes")
        sheet.append(header)
        overlap = Counter()
        appended = 0
        try:
            if baseline:
                with report_rows(baseline) as (_, rows):
                    for row in rows:
                        day = parse_date(row[date_index])
                        sheet.append(row)
                        if day >= start_day:
                            overlap[_row_key(row, date_index)] += 1
            with report_rows(incoming) as (_, rows):
                for row in rows:
                    key = _row_key(row, date_index)
                    if overlap[key]:
                        overlap[key] -= 1
                    else:
                        sheet.append(row)
                        appended += 1
            checkpoint = workbook.create_sheet(CHECKPOINT_SHEET)
            checkpoint.sheet_state = "hidden"
            checkpoint.append(["last_successful_day", end_day.isoformat()])
            checkpoint.append(["updated_at", datetime.now(LIMA_TZ).isoformat()])
            workbook.save(staged)
        finally:
            workbook.close()
        os.replace(staged, target)

    # These are old full snapshots generated by the previous downloader.
    # Never delete project-root seed files or unrelated spreadsheets.
    for old_file in legacy_reports(directory):
        try:
            old_file.unlink()
        except OSError as error:
            print(f"Could not remove old canjes snapshot {old_file.name}: {error}")
    print(f"Appended {appended} canjes; checkpoint: {end_day.isoformat()}")
    return str(target)
